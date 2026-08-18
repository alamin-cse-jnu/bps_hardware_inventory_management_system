"""
Tests for the Office-wise Asset List report (Phase 11).

Covers the three pieces that can silently go wrong:
  * scope resolution — especially the narrowing rule and the "only" flags
  * grouping / merge-run computation
  * Excel merge geometry, checked against the shape of
    docs/office wise asset list.xlsx
"""

from django.contrib.auth import get_user_model
from django.test import Client, RequestFactory, TestCase
from django.urls import reverse

from assets.models import AssetCategory, AssetItem, AssetType
from assignees.models import Assignee, AssigneeType, CachedEmployee, CachedOffice
from assignments.models import Assignment

User = get_user_model()

XLSX_MAGIC = b"PK"

_BRANCH_NAMES = {"110": "Alpha Branch One", "120": "Alpha Branch Two", "210": "Beta Branch"}
_SECTION_NAMES = {"111": "Section A", "112": "Section B"}
_WING_NAMES = {"100": "Alpha Wing", "200": "Beta Wing"}


class OfficeFixture(TestCase):
    """
    Miniature Secretariat:

        ROOT
         └── SECRETARY
              ├── Alpha Wing
              │    ├── Alpha Branch One
              │    │    ├── Section A   (└── Unit U)
              │    │    └── Section B
              │    └── Alpha Branch Two
              └── Beta Wing
                   └── Beta Branch
    """

    def setUp(self):
        self.user = User.objects.create_user(
            username="office-tester", password="testpass123!",
            is_superuser=True, is_staff=True,
        )
        self.offices = {}
        for prp, parent, name in [
            ("1",   "0",   "ROOT"),
            ("10",  "1",   "SECRETARY"),
            ("100", "10",  "Alpha Wing"),
            ("110", "100", "Alpha Branch One"),
            ("111", "110", "Section A"),
            ("112", "110", "Section B"),
            ("113", "111", "Unit U"),
            ("120", "100", "Alpha Branch Two"),
            ("200", "10",  "Beta Wing"),
            ("210", "200", "Beta Branch"),
        ]:
            self.offices[name] = CachedOffice.objects.create(
                prp_id=prp, parent_prp_id=parent, name_en=name
            )

        self.cat_computing, _ = AssetCategory.objects.get_or_create(name="Computing Equipment")
        self.cat_network, _ = AssetCategory.objects.get_or_create(name="Networking Equipment")
        self.t_laptop, _ = AssetType.objects.get_or_create(
            category=self.cat_computing, name="Laptop", defaults={"spec_schema": []})
        self.t_monitor, _ = AssetType.objects.get_or_create(
            category=self.cat_computing, name="Monitor", defaults={"spec_schema": []})
        self.t_ap, _ = AssetType.objects.get_or_create(
            category=self.cat_network, name="Access Point", defaults={"spec_schema": []})

    # ── fixture helpers ──────────────────────────────────────────────────────

    def employee(self, name, wing="", branch="", section="", unit="",
                 designation="Officer", is_active=True):
        emp = CachedEmployee.objects.create(
            name_en=name, source="MANUAL", designation_en=designation,
            is_active=is_active,
            wing_id=wing, wing_name_en=_WING_NAMES.get(wing, ""),
            branch_id=branch, branch_name_en=_BRANCH_NAMES.get(branch, ""),
            section_id=section, section_name_en=_SECTION_NAMES.get(section, ""),
            unit_id=unit, unit_name_en="Unit U" if unit else "",
        )
        return Assignee.objects.create(assignee_type=AssigneeType.EMPLOYEE, employee=emp)

    def office_assignee(self, name):
        return Assignee.objects.create(
            assignee_type=AssigneeType.OFFICE, office=self.offices[name]
        )

    def give(self, assignee, tag, asset_type):
        asset = AssetItem.objects.create(
            asset_tag=tag, asset_type=asset_type, brand="Acme", model_name="M1",
            serial_number=f"SN-{tag}",
            status=AssetItem.Status.ASSIGNED, created_by=self.user,
        )
        Assignment.objects.create(
            asset=asset, assignee=assignee, performed_by=self.user,
            holder_snapshot=assignee.build_snapshot(),
        )
        return asset

    def scope_count(self, **kwargs):
        from reports.office_scope import (
            OfficeScope, build_office_options, scope_terms, terms_to_q,
        )
        options = build_office_options()
        terms = scope_terms(OfficeScope(**kwargs), options)
        return CachedEmployee.objects.filter(terms_to_q(terms)).count()


# ── Scope resolution ─────────────────────────────────────────────────────────

class OfficeScopeLogicTest(OfficeFixture):
    def setUp(self):
        super().setUp()
        self.employee("Wing Level", wing="100")
        self.employee("Branch Level", wing="100", branch="110")
        self.employee("Section Person", wing="100", branch="110", section="111")
        self.employee("Unit Person", wing="100", branch="110", section="111", unit="113")
        self.employee("Other Branch", wing="100", branch="120")
        self.employee("Beta Person", wing="200", branch="210")

    def test_options_derived_from_tree(self):
        from reports.office_scope import build_office_options
        opt = build_office_options()
        self.assertEqual([w["name"] for w in opt.wings], ["Alpha Wing", "Beta Wing"])
        self.assertEqual([b["name"] for b in opt.branches],
                         ["Alpha Branch One", "Alpha Branch Two", "Beta Branch"])
        self.assertEqual([s["name"] for s in opt.sections], ["Section A", "Section B"])

    def test_empty_scope_covers_all_wings(self):
        self.assertEqual(self.scope_count(), 6)

    def test_wing_includes_descendants(self):
        self.assertEqual(self.scope_count(wing=["100"]), 5)

    def test_wing_only_is_that_office_alone(self):
        self.assertEqual(self.scope_count(wing=["100"], wing_only=True), 1)

    def test_multiple_wings_union(self):
        self.assertEqual(self.scope_count(wing=["100", "200"]), 6)

    def test_branch_narrows_within_selected_wing(self):
        """Picking a branch must not be swallowed by its parent wing's term."""
        self.assertEqual(self.scope_count(wing=["100"], branch=["110"]), 3)

    def test_branch_only(self):
        self.assertEqual(self.scope_count(branch=["110"], branch_only=True), 1)

    def test_wing_only_plus_branch_adds_both(self):
        """1 wing-level person + the 3 in the branch subtree."""
        self.assertEqual(
            self.scope_count(wing=["100"], wing_only=True, branch=["110"]), 4
        )

    def test_section_narrows_within_branch(self):
        self.assertEqual(
            self.scope_count(wing=["100"], branch=["110"], section=["111"]), 2
        )

    def test_section_only_excludes_units(self):
        self.assertEqual(self.scope_count(section=["111"], section_only=True), 1)

    def test_unknown_ids_are_discarded(self):
        from reports.office_scope import build_office_options, parse_scope
        request = RequestFactory().get("/", {"wing": ["100", "99999"]})
        self.assertEqual(parse_scope(request, build_office_options()).wing, ["100"])

    def test_comma_separated_ids_accepted(self):
        from reports.office_scope import build_office_options, parse_scope
        request = RequestFactory().get("/?wing=100,200")
        self.assertEqual(parse_scope(request, build_office_options()).wing, ["100", "200"])


# ── Grouping and merge runs ──────────────────────────────────────────────────

class OfficeAssetGroupingTest(OfficeFixture):
    def setUp(self):
        super().setUp()
        from reports.office_scope import OfficeScope, build_office_options, scope_terms

        wing_person = self.employee("Wing Boss", wing="100", designation="Director General")
        sec_person = self.employee("Section Staff", wing="100", branch="110", section="111")

        self.give(wing_person, "AP-001", self.t_ap)
        self.give(sec_person, "LAP-001", self.t_laptop)
        self.give(sec_person, "LAP-002", self.t_laptop)
        self.give(sec_person, "MON-001", self.t_monitor)

        self.options = build_office_options()
        self.terms = scope_terms(OfficeScope(), self.options)

    def groups(self):
        from reports.office_assets import build_groups
        return build_groups(self.terms, self.options)

    def test_grouped_by_holder(self):
        groups = self.groups()
        self.assertEqual([g["holder"] for g in groups], ["Wing Boss", "Section Staff"])
        self.assertEqual([g["rowspan"] for g in groups], [1, 3])

    def test_missing_levels_render_as_ellipsis(self):
        from reports.office_assets import MISSING
        boss = self.groups()[0]
        self.assertEqual(boss["wing"], "Alpha Wing")
        self.assertEqual(boss["branch"], MISSING)
        self.assertEqual(boss["section"], MISSING)

    def test_run_spans_within_holder(self):
        rows = self.groups()[1]["rows"]
        self.assertEqual([r["asset_tag"] for r in rows], ["LAP-001", "LAP-002", "MON-001"])
        # All three are Computing Equipment → a single run of 3.
        self.assertEqual([r["_span_category"] for r in rows], [3, 0, 0])
        # Laptop, Laptop, Monitor → runs of 2 then 1.
        self.assertEqual([r["_span_asset_type"] for r in rows], [2, 0, 1])

    def test_spans_never_cross_holders(self):
        for group in self.groups():
            self.assertEqual(
                sum(r["_span_category"] for r in group["rows"]), len(group["rows"])
            )

    def test_inactive_holder_is_flagged_not_dropped(self):
        ghost = self.employee("Gone Person", wing="100", branch="110", is_active=False)
        self.give(ghost, "LAP-999", self.t_laptop)
        match = [g for g in self.groups() if g["holder"].startswith("Gone Person")]
        self.assertEqual(len(match), 1)
        self.assertTrue(match[0]["_inactive"])
        self.assertIn("Inactive", match[0]["holder"])

    def test_office_held_assets_sort_after_that_office_staff(self):
        self.give(self.office_assignee("Alpha Wing"), "PRI-001", self.t_laptop)
        groups = self.groups()
        names = [g["holder"] for g in groups]
        self.assertEqual(names.index("Wing Boss") + 1, names.index("Alpha Wing"))
        self.assertLess(names.index("Alpha Wing"), names.index("Section Staff"))
        office_group = groups[names.index("Alpha Wing")]
        self.assertTrue(office_group["_is_office"])
        self.assertEqual(office_group["wing"], "Alpha Wing")

    def test_office_outside_scope_excluded(self):
        from reports.office_assets import build_groups
        from reports.office_scope import OfficeScope, scope_terms

        self.give(self.office_assignee("Beta Wing"), "PRI-002", self.t_laptop)
        terms = scope_terms(OfficeScope(wing=["100"]), self.options)
        self.assertNotIn("Beta Wing", [g["holder"] for g in build_groups(terms, self.options)])

    def test_returned_assignments_excluded(self):
        from django.utils import timezone
        Assignment.objects.filter(asset__asset_tag="LAP-001").update(
            returned_at=timezone.now()
        )
        tags = [r["asset_tag"] for g in self.groups() for r in g["rows"]]
        self.assertNotIn("LAP-001", tags)


# ── Excel merge geometry ─────────────────────────────────────────────────────

class OfficeAssetsExcelTest(OfficeAssetGroupingTest):
    def sheet(self):
        import io

        import openpyxl

        from reports.generators.excel import office_assets_excel
        data = office_assets_excel(self.groups(), subtitle="Test scope")
        self.assertTrue(data.startswith(XLSX_MAGIC))
        return openpyxl.load_workbook(io.BytesIO(data)).active

    def test_merge_ranges_match_reference_shape(self):
        merges = {str(m) for m in self.sheet().merged_cells.ranges}
        # Title rows 1-3, header row 5, data from row 6. Column A is SL, so
        # the reference layout's A-E identity block lives in B-F here.
        # Wing Boss = row 6 (single row); Section Staff = rows 7-9.
        self.assertIn("A7:A9", merges)      # SL merged across the holder
        for col in "BCDEF":
            self.assertIn(f"{col}7:{col}9", merges, f"identity column {col} not merged")
        self.assertIn("G7:G9", merges)      # Category across the whole block
        self.assertIn("H7:H8", merges)      # Asset Type only across the 2 laptops
        self.assertNotIn("H7:H9", merges)
        for col in ("I", "J", "K", "L"):    # Brand / Model / Serial / Tag never merge
            self.assertFalse(any(m.startswith(col) for m in merges),
                             f"column {col} must not be merged")

    def test_single_row_holder_has_no_identity_merge(self):
        merges = {str(m) for m in self.sheet().merged_cells.ranges}
        self.assertNotIn("B6:B6", merges)
        self.assertNotIn("A6:A6", merges)

    def test_sl_numbers_holders_not_rows(self):
        ws = self.sheet()
        self.assertEqual(ws["A6"].value, 1)   # Wing Boss
        self.assertEqual(ws["A7"].value, 2)   # Section Staff, 3 assets
        self.assertIsNone(ws["A8"].value)     # swallowed by the merge
        self.assertIsNone(ws["A9"].value)

    def test_continuation_cells_are_blank(self):
        ws = self.sheet()
        self.assertIsNone(ws["G8"].value)
        self.assertIsNone(ws["H8"].value)
        self.assertEqual(ws["L8"].value, "LAP-002")

    def test_serial_number_precedes_asset_tag(self):
        ws = self.sheet()
        self.assertEqual(ws["K7"].value, "SN-LAP-001")
        self.assertEqual(ws["L7"].value, "LAP-001")

    def test_identity_values_written_once_at_anchor(self):
        ws = self.sheet()
        self.assertEqual(ws["B7"].value, "Section Staff")
        self.assertIsNone(ws["B8"].value)
        self.assertIsNone(ws["B9"].value)

    def test_all_cells_bordered(self):
        """
        Checked against the written XML, not a re-read workbook: openpyxl's
        reader reports every MergedCell's style as default regardless of what
        the file actually contains, so a round-trip read cannot see borders on
        the non-anchor cells of a merged block.
        """
        import io
        import zipfile
        from xml.etree import ElementTree as ET

        from reports.generators.excel import office_assets_excel

        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        archive = zipfile.ZipFile(io.BytesIO(office_assets_excel(self.groups())))
        styles = ET.fromstring(archive.read("xl/styles.xml"))

        # Border definitions that carry a thin line on all four edges.
        fully_bordered = {
            i for i, border in enumerate(styles.find(f"{ns}borders"))
            if all(
                (edge := border.find(f"{ns}{name}")) is not None
                and edge.get("style") == "thin"
                for name in ("left", "right", "top", "bottom")
            )
        }
        self.assertTrue(fully_bordered, "no fully-bordered style was defined")

        xf_border = [
            int(xf.get("borderId", 0)) for xf in styles.find(f"{ns}cellXfs")
        ]
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        style_of = {
            c.get("r"): int(c.get("s", 0)) for c in sheet.iter(f"{ns}c")
        }

        # Header row 5 through the last data row 9, all 12 columns.
        for r in range(5, 10):
            for col in "ABCDEFGHIJKL":
                coord = f"{col}{r}"
                self.assertIn(coord, style_of, f"{coord} not written")
                self.assertIn(
                    xf_border[style_of[coord]], fully_bordered,
                    f"{coord} is not fully bordered",
                )

    def test_header_labels(self):
        ws = self.sheet()
        self.assertEqual(
            [ws.cell(row=5, column=c).value for c in range(1, 13)],
            ["SL", "Holder", "Designation", "Wing", "Branch", "Section",
             "Category", "Asset Type", "Brand", "Model", "Serial Number",
             "Asset Tag"],
        )

    def test_empty_groups_still_produce_a_workbook(self):
        from reports.generators.excel import office_assets_excel
        self.assertTrue(office_assets_excel([]).startswith(XLSX_MAGIC))


# ── Views ────────────────────────────────────────────────────────────────────

class OfficeAssetsViewTest(OfficeFixture):
    def setUp(self):
        super().setUp()
        self.client = Client()
        self.client.force_login(self.user)
        person = self.employee("Viewer Target", wing="100", branch="110", section="111")
        self.give(person, "LAP-100", self.t_laptop)
        self.give(person, "LAP-101", self.t_laptop)

    def test_view_page_renders(self):
        resp = self.client.get(reverse("reports:office_assets_view"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Office-wise Asset List")
        self.assertContains(resp, "Viewer Target")

    def test_view_page_emits_rowspan(self):
        resp = self.client.get(reverse("reports:office_assets_view"))
        self.assertContains(resp, 'rowspan="2"')

    def test_scope_filter_applies(self):
        resp = self.client.get(reverse("reports:office_assets_view") + "?wing=200")
        self.assertNotContains(resp, "Viewer Target")

    def test_wing_only_excludes_deeper_staff(self):
        resp = self.client.get(
            reverse("reports:office_assets_view") + "?wing=100&wing_only=1")
        self.assertNotContains(resp, "Viewer Target")

    def test_excel_download(self):
        resp = self.client.get(reverse("reports:office_assets_excel") + "?wing=100")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn("office_wise_asset_list.xlsx", resp["Content-Disposition"])
        self.assertTrue(resp.content.startswith(XLSX_MAGIC))

    def test_counts_in_context(self):
        resp = self.client.get(reverse("reports:office_assets_view") + "?per_page=25")
        self.assertEqual(resp.context["holder_count"], 1)
        self.assertEqual(resp.context["asset_total"], 2)

    def test_requires_login(self):
        self.client.logout()
        resp = self.client.get(reverse("reports:office_assets_view"))
        self.assertIn(resp.status_code, (302, 403))
